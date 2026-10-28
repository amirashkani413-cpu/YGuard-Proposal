#!/usr/bin/env python3
"""Generate editable Word and print-ready PDF from proposal-fa.md."""
from pathlib import Path
import re, math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Flowable, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "files"
OUT.mkdir(exist_ok=True)
text = (ROOT / "proposal-fa.md").read_text(encoding="utf-8")
DOCX = OUT / "YGuard_Club_Partnership_Proposal_FA.docx"
PDF = OUT / "YGuard_Club_Partnership_Proposal_FA.pdf"
REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
if not Path(REG).exists():
    raise RuntimeError("DejaVu Sans font missing")
def plain(s):
    return re.sub(r"[*`]", "", s).strip()
def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = p._p.get_or_add_pPr()
    el = OxmlElement("w:bidi")
    el.set(qn("w:val"), "1")
    ppr.append(el)
    for run in p.runs:
        rpr = run._r.get_or_add_rPr()
        el = OxmlElement("w:rtl")
        el.set(qn("w:val"), "1")
        rpr.append(el)
        run.font.name = "Arial"
        run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

# Parse markdown into blocks, preserving the pricing table.
blocks = []
lines = text.splitlines()
i = 0
while i < len(lines):
    s = lines[i].strip()
    if not s or s == "---":
        i += 1; continue
    if s.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cols = [plain(x) for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[:\- ]+", c) for c in cols):
                rows.append(cols)
            i += 1
        blocks.append(("table", rows)); continue
    if s.startswith("# "): kind, val = "title", s[2:]
    elif s.startswith("## "): kind, val = "heading", s[3:]
    elif s.startswith("- "): kind, val = "bullet", s[2:]
    elif re.match(r"^\d+\. ", s): kind, val = "bullet", s
    else: kind, val = "paragraph", s
    blocks.append((kind, plain(val)))
    i += 1

# Editable Word
doc = Document()
sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.65), Cm(1.65)
style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(9.5)
style.font.color.rgb = RGBColor(35, 49, 64)
style.paragraph_format.space_after = Pt(3)

for kind, val in blocks:
    if kind == "table":
        tbl = doc.add_table(rows=0, cols=len(val[0]))
        tbl.style = "Table Grid"
        for j, row in enumerate(val):
            cells = tbl.add_row().cells
            for k, item in enumerate(row):
                cells[k].text = item
                if j == 0: shade(cells[k], "183A56")
                for p in cells[k].paragraphs:
                    rtl(p)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(8)
                        if j == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(255,255,255)
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3 if kind != "heading" else 4)
    if kind == "heading":
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.keep_with_next = True
    if kind == "title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
    if kind == "bullet":
        p.paragraph_format.right_indent = Cm(.25)
        val = "•  " + val
    r = p.add_run(val)
    if kind in ("title", "heading"):
        r.bold = True
        r.font.size = Pt(16 if kind == "title" else 11.5)
        r.font.color.rgb = RGBColor(20, 58, 88)
    elif kind == "bullet":
        r.font.size = Pt(9.1)
    rtl(p)
    if kind == "title": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("YGuard  |  Yguard.ir")
r.font.size = Pt(8)
doc.save(DOCX)

# PDF uses Pillow's libraqm for correct Arabic/Persian shaping and ReportLab Platypus
# for automatic page breaking. Fonts are used only inside the generated documents.
SCALE = 3
class RTL(Flowable):
    def __init__(self, txt, width, size=9, lead=13, bold=False, color="#233140", center=False, after=0):
        Flowable.__init__(self)
        self.txt, self.width, self.size, self.lead = txt, width, size, lead
        self.color, self.center, self.after = color, center, after
        self.font = ImageFont.truetype(BOLD if bold else REG, int(size*SCALE))
        self.lines = []
    def wrap(self, aw, ah):
        self.width = min(self.width, aw)
        maxlen = (self.width-3)*SCALE
        current = ""
        for word in self.txt.split():
            candidate = (current+" "+word).strip()
            if self.font.getlength(candidate, direction="rtl", language="fa") <= maxlen:
                current = candidate
            else:
                if current: self.lines.append(current)
                current = word
        if current: self.lines.append(current)
        if not self.lines: self.lines = [""]
        self.height = len(self.lines)*self.lead + self.after
        return self.width, self.height
    def draw(self):
        im = Image.new("RGBA", (math.ceil(self.width*SCALE), math.ceil(self.height*SCALE)), (255,255,255,0))
        d = ImageDraw.Draw(im)
        for i, line in enumerate(self.lines):
            y = int(i*self.lead*SCALE)
            if self.center:
                d.text((im.width//2, y), line, font=self.font, fill=self.color, anchor="mt", direction="rtl", language="fa")
            else:
                d.text((im.width-4, y), line, font=self.font, fill=self.color, anchor="rt", direction="rtl", language="fa")
        self.canv.drawImage(ImageReader(im), 0, 0, width=self.width, height=self.height, mask="auto")

W, H = A4
width = W - 30*mm
story = []
for kind, val in blocks:
    if kind == "title":
        story.append(RTL("YGUARD", width, size=23, lead=31, bold=True, color="#143A58", center=True, after=4))
        story.append(RTL(val, width, size=15, lead=21, bold=True, color="#143A58", center=True, after=10))
    elif kind == "heading":
        story.append(Spacer(1, 4*mm))
        story.append(RTL(val, width, size=11, lead=17, bold=True, color="#143A58", after=2))
    elif kind == "table":
        cw = [width*.18, width*.24, width*.12, width*.29, width*.17]
        data = []
        for rownum, row in enumerate(val):
            data.append([RTL(row[j], cw[j]-5, size=7.3, lead=11, bold=(rownum==0),
                             color="#FFFFFF" if rownum==0 else "#233140", center=True)
                         for j in range(len(cw))])
        tbl = Table(data, colWidths=cw, repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#183A56")),
            ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#D5DDE5")),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("TOPPADDING", (0,0), (-1,-1), 5),
            ("BOTTOMPADDING", (0,0), (-1,-1), 5),
            ("LEFTPADDING", (0,0), (-1,-1), 2),
            ("RIGHTPADDING", (0,0), (-1,-1), 2),
        ]))
        story.append(tbl)
    else:
        story.append(RTL(("•  " if kind=="bullet" else "")+val,
                         width-3*mm, size=8.5 if kind=="bullet" else 9,
                         lead=12.6 if kind=="bullet" else 13.5,
                         after=2.5, color="#233140"))

def footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D5DDE5"))
    canvas.line(15*mm, 13*mm, W-15*mm, 13*mm)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#718096"))
    canvas.drawCentredString(W/2, 8*mm, "YGuard  |  Yguard.ir")
    canvas.drawRightString(W-15*mm, 8*mm, str(doc.page))
    canvas.restoreState()

pdf = SimpleDocTemplate(str(PDF), pagesize=A4, leftMargin=15*mm,
                        rightMargin=15*mm, topMargin=12*mm, bottomMargin=18*mm,
                        title="YGuard Partnership Proposal")
pdf.build(story, onFirstPage=footer, onLaterPages=footer)
print("Generated:", DOCX, PDF)
